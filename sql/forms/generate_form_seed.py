#!/usr/bin/env python3
"""
generate_form_seed.py - Forms engine seed generator (PM build slice 1).

Reads the two workbooks in sql/forms/source/ and writes three T-SQL seed scripts next to
this file so the workbook stays the source of truth:

  seed_form_template_hvac.sql       FormAnswerList + FormTemplate/Section/Question for HVAC
                                    (Master Data Fields + PM Tasks + Photo Requirements)
  seed_asset_categories_hvac.sql    the Dropdown Lists sheet's 12 equipment types as HVAC AssetCategory rows
  seed_form_rules_pm_hvac.sql       one FormRule + PMRule per row of the PM Pricing Info sheet

Plus seed_form_report.txt: counts and every row-level judgement the generator made
(skipped rows, unstructured triggers, name aliases). The human-readable rationale lives in
Temp\PM\technical-design\SEED_DECISIONS.md; keep the two in sync when rules change here.

Seed rules follow plan section 8a (rules 1-14). Search for "rule N" below.

Usage:  python generate_form_seed.py [--dict PATH] [--pricing PATH] [--out DIR]
Requires openpyxl.
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import json
import re
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
DEFAULT_DICT = HERE / "source" / "Universal_HVAC_PM_Master_Checklist_Data_Dictionary.xlsx"
DEFAULT_PRICING = HERE / "source" / "Copy of PM Pricing Info.xlsx"

TEMPLATE_NAME = "HVAC PM"
PARENT_TRADE = "HVAC"          # resolved to t_id at run time (t_id_parent IS NULL)
TODAY = dt.date.today().isoformat()

# CheckListAnswerType rows reused as-is (no rows are added - see create_form_tables.sql header)
CLAT_TEXTBOX, CLAT_DROPDOWN, CLAT_SIGNATURE, CLAT_PHOTO, CLAT_READONLY = 1, 2, 3, 4, 5

# ---------------------------------------------------------------------------------------
# Vocabularies (also documented in SEED_DECISIONS.md)
# ---------------------------------------------------------------------------------------
EQUIPMENT_TYPES = ["RTU", "Split System", "AHU", "Heat Pump", "Water-Source Heat Pump", "VAV", "FPVAV",
                   "Fan Coil", "VRF", "Exhaust / Booster Fan", "Refrigeration", "Other"]          # rule 13
COOLING = ["RTU", "Split System", "AHU", "Heat Pump", "Water-Source Heat Pump", "Fan Coil", "VRF", "Refrigeration", "Other"]
REFRIG = ["RTU", "Split System", "Heat Pump", "Water-Source Heat Pump", "VRF", "Refrigeration"]     # compressor-bearing
CONDFAN = ["RTU", "Split System", "Heat Pump", "VRF", "Refrigeration"]                             # air-cooled condenser fans
HEATING_TYPES = ["Gas", "Electric", "Heat Pump", "Hot Water", "Steam", "Oil", "None", "Other"]     # F059 (superset of the sheet's 6)
HEATING_ANY = [h for h in HEATING_TYPES if h != "None"]
FUEL_FIRED = ["Gas", "Oil"]
HYDRONIC = {"any": [{"heatingType": ["Hot Water", "Steam"]}, {"equipmentType": ["Water-Source Heat Pump"]}]}
ROOF_MOUNTS = ["Rooftop", "Stilts", "Fenced area on the roof"]

SEASON_COOL = "Spring Cooling;Full PM"
SEASON_HEAT = "Fall Heating;Full PM"

# Section catalogue: (name, phase, repeat per unit, condition JSON)   -- rule 5 + plan 5.2
SECTIONS = [
    ("Site & Arrival", "Site", 0, None),
    ("Asset: Identification", "Asset", 1, None),
    ("Asset: Filters & Belts", "Asset", 1, None),
    ("Asset: Controls & Features", "Asset", 1, None),
    ("Cabinet & Service Area", "Unit", 1, None),
    ("Filters, Belts, Fans & Airflow", "Unit", 1, None),
    ("Coils & Condensate", "Unit", 1, {"equipmentType": COOLING}),
    ("Electrical, Thermostat & Controls", "Unit", 1, None),
    ("Cooling & Refrigeration Readings", "Unit", 1, {"equipmentType": COOLING}),
    ("Heating & Hydronic Readings", "Unit", 1, {"any": [{"heatingType": HEATING_ANY}, {"equipmentType": ["Water-Source Heat Pump"]}]}),
    ("Unit Condition", "Unit", 1, None),
    ("Recommendations / Quote Readiness", "Unit", 1, None),
    ("Closeout & Signoff", "Checkout", 0, None),
]
SECTION_INDEX = {name: i for i, (name, *_rest) in enumerate(SECTIONS)}

# Master Data Fields: (section number, Subsection) -> section name
FIELD_SECTION_MAP = {
    (4, "Work Order & Arrival"): "Site & Arrival",
    (4, "Filters, Belts, Fans & Airflow"): "Site & Arrival",          # F101 indoor temp, F102 outdoor temp
    (5, "Identification & Condition"): "Asset: Identification",
    (5, "Filters, Belts, Fans & Airflow"): "Asset: Filters & Belts",
    (5, "Electrical, Thermostat & Controls"): "Asset: Controls & Features",
    (6, "Identification & Condition"): "Unit Condition",
    (6, "Filters, Belts, Fans & Airflow"): "Filters, Belts, Fans & Airflow",
    (6, "Cooling & Refrigeration Readings"): "Cooling & Refrigeration Readings",
    (6, "Electrical, Thermostat & Controls"): "Electrical, Thermostat & Controls",
    (6, "Heating & Hydronic Readings"): "Heating & Hydronic Readings",
    (6, "Recommendations / Quote Readiness"): "Recommendations / Quote Readiness",
    (7, "Work Order & Arrival"): "Closeout & Signoff",
    (7, "Completion & Signoff"): "Closeout & Signoff",
}
PHASE_BY_SECTION_NUMBER = {4: "Site", 5: "Asset", 6: "Unit", 7: "Checkout"}                       # rule 5

# PM Tasks: Equipment Area -> section name (phase from the section)                              -- rule 5
TASK_AREA_MAP = {
    "Arrival / Site": "Site & Arrival",
    "Cabinet / Exterior": "Cabinet & Service Area",
    "Roof / Service Area": "Cabinet & Service Area",
    "Filters": "Filters, Belts, Fans & Airflow",
    "Fresh Air": "Filters, Belts, Fans & Airflow",
    "Fans / Motors": "Filters, Belts, Fans & Airflow",
    "Belts / Pulleys": "Filters, Belts, Fans & Airflow",
    "Air Distribution": "Filters, Belts, Fans & Airflow",
    "Performance": "Filters, Belts, Fans & Airflow",
    "Coils": "Coils & Condensate",
    "Condensate": "Coils & Condensate",
    "Heating Coil": "Heating & Hydronic Readings",
    "Electrical": "Electrical, Thermostat & Controls",
    "Thermostat / Controls": "Electrical, Thermostat & Controls",
    "EMS / DDC": "Electrical, Thermostat & Controls",
    "Economizer": "Electrical, Thermostat & Controls",
    "Refrigeration": "Cooling & Refrigeration Readings",
    "Compressor": "Cooling & Refrigeration Readings",
    "Condenser Fans": "Cooling & Refrigeration Readings",
    "Heat Pump": "Heating & Hydronic Readings",
    "Gas Heating": "Heating & Hydronic Readings",
    "Electric Heating": "Heating & Hydronic Readings",
    "Hydronic Heating": "Heating & Hydronic Readings",
    "Hydronic / Water Cooled": "Heating & Hydronic Readings",
    "Hydronic / Seasonal": "Heating & Hydronic Readings",
    "Heating Performance": "Heating & Hydronic Readings",
    "Closeout": "Closeout & Signoff",
}

# Photo Requirements: Photo ID -> (section, repeat key)                                          -- rule 8
PHOTO_SECTION_MAP = {
    "P01": ("Site & Arrival", None), "P02": ("Cabinet & Service Area", None), "P03": ("Cabinet & Service Area", None),
    "P04": ("Asset: Identification", None), "P05": ("Asset: Identification", None), "P06": ("Asset: Identification", None),
    "P07": ("Filters, Belts, Fans & Airflow", "FilterBank"), "P08": ("Filters, Belts, Fans & Airflow", "FilterBank"),
    "P09": ("Electrical, Thermostat & Controls", None),
    "P10": ("Coils & Condensate", None), "P11": ("Coils & Condensate", None), "P12": ("Coils & Condensate", None), "P13": ("Coils & Condensate", None),
    "P14": ("Filters, Belts, Fans & Airflow", "BeltDrive"), "P15": ("Filters, Belts, Fans & Airflow", None),
    "P16": ("Electrical, Thermostat & Controls", None), "P17": ("Electrical, Thermostat & Controls", None), "P18": ("Electrical, Thermostat & Controls", None),
    "P19": ("Filters, Belts, Fans & Airflow", None), "P20": ("Cooling & Refrigeration Readings", "Circuit"),
    "P21": ("Electrical, Thermostat & Controls", None), "P22": ("Heating & Hydronic Readings", None), "P23": ("Heating & Hydronic Readings", None),
    "P24": ("Cabinet & Service Area", None), "P25": ("Recommendations / Quote Readiness", None),
    "P26": ("Cabinet & Service Area", None), "P27": ("Cabinet & Service Area", None), "P28": ("Closeout & Signoff", None),
}
# Photo trigger -> (photorequired, requirement, condition, linkedcode)
PHOTO_TRIGGER_MAP = {
    "P01": ("CustomerConfigured", "Configured", None, "T001"),
    "P02": ("Recommended", "Recommended", None, None),
    "P03": ("Always", "Always", None, None),
    "P04": ("Always", "Always", None, None),
    "P05": ("Always", "Always", None, None),
    "P06": ("Always", "Conditional", {"equipmentType": ["Split System", "AHU"]}, None),
    "P07": ("CustomerConfigured", "Configured", {"whenCode": "F076", "whenIn": ["Yes"]}, "T012"),          # rule 7: photo stamp group
    "P08": ("CustomerConfigured", "Configured", {"whenCode": "F076", "whenIn": ["Yes"]}, "T013"),
    "P09": ("CustomerConfigured", "Configured", None, None),
    "P10": ("Always", "Conditional", {"whenCode": "T026", "whenIn": ["Completed", "Deep clean quote"]}, "T024"),
    "P11": ("Always", "Conditional", {"whenCode": "T026", "whenIn": ["Completed", "Deep clean quote"]}, "T025"),
    "P12": ("CustomerConfigured", "Configured", {"equipmentType": COOLING}, "T030"),
    "P13": ("CustomerConfigured", "Configured", None, "T033"),
    "P14": ("Recommended", "Conditional", {"attr": {"BeltDriven": True}}, "T021"),
    "P15": ("IfIssue", "Conditional", None, "T018"),
    "P16": ("CustomerConfigured", "Configured", None, "T043"),
    "P17": ("CustomerConfigured", "Configured", None, "T044"),
    "P18": ("IfIssue", "Conditional", None, "T047"),
    "P19": ("CustomerConfigured", "Configured", None, "T074"),
    "P20": ("Recommended", "Conditional", {"equipmentType": REFRIG}, "T054"),
    "P21": ("IfIssue", "Conditional", None, "T039"),
    "P22": ("Recommended", "Conditional", {"heatingType": FUEL_FIRED}, "T065"),
    "P23": ("Recommended", "Conditional", {"heatingType": FUEL_FIRED}, "T067"),
    "P24": ("IfIssue", "Conditional", None, "T009"),
    "P25": ("IfIssue", "Conditional", None, None),          # any deficiency on the unit
    "P26": ("Recommended", "Recommended", None, None),
    "P27": ("CustomerConfigured", "Configured", None, "T011"),
    "P28": ("CustomerConfigured", "Configured", None, "T083"),
}
PHOTO_TIMING_MAP = {
    "arrival": "Before", "before work": "Before", "after work": "After", "before/after": "BeforeAfter",
    "before and after cleaning": "BeforeAfter", "before/after cleaning": "BeforeAfter", "before/after adjustment": "BeforeAfter",
    "during test": "During", "during operation": "During", "during inspection": "During", "when found": "During", "closeout": "After",
}

# Configured questions that a PM rule column turns on (rule 7, plus additions logged in SEED_DECISIONS.md)
CONFIGURED_BY_RULE = {
    "setpoints": ["F135", "F136", "F137", "F139", "F140", "F141", "F142", "T043", "T044", "T045"],
    "antialgae": ["T033", "P13"],
    "phototimestamp": ["P07", "P08", "P09", "P12", "P16", "P17"],
    "managerseesoldfilters": ["F081", "T015"],
    "signoff": ["T083", "P28"],
    "ivrrequired": ["T084"],
    "filtersincluded": ["T013", "T014"],
    "immediatequoterequired": ["T086"],
}
FORCE_CONFIGURED = {code for codes in CONFIGURED_BY_RULE.values() for code in codes}

# Trigger / Applies When text -> structured condition (rule 12). Keys are lowercase.
TRIGGER_CONDITIONS = {
    # asset facts
    "belt-driven equipment": {"attr": {"BeltDriven": True}},
    "economizer present": {"attr": {"Economizer": True}},
    "economizer/damper present": {"attr": {"Economizer": True}},
    "economizer/mixing section present": {"attr": {"Economizer": True}},
    "outdoor-air system present": {"attr": {"Economizer": True}},
    "outdoor-air intake present": {"attr": {"Economizer": True}},
    "ems/bas present": {"attr": {"EMS": True}},
    "ddc/ems present": {"attr": {"EMS": True}},
    "remote sensors present": {"attr": {"RemoteSensors": True}},
    "thermostat controls remote sensors": {"attr": {"RemoteSensors": True}},
    "three-phase equipment": {"attr": {"ThreePhase": True}},
    "ducted equipment": {"attr": {"Ducted": True}},
    "ducted equipment / scope requires": {"attr": {"Ducted": True}},
    "accessible duct system": {"attr": {"Ducted": True}},
    "fresh-air filters present": {"attr": {"FreshAirFilter": True}},
    "fresh-air filter present": {"attr": {"FreshAirFilter": True}},
    "crankcase heater present / fall scope": {"attr": {"CrankcaseHeater": True}},
    "glycol system": {"attr": {"Glycol": True}},
    "battery thermostat": {"attr": {"BatteryThermostat": True}},
    "battery thermostat / batteries weak": {"attr": {"BatteryThermostat": True}},
    "split systems or connected assets": {"equipmentType": ["Split System", "AHU", "VRF"]},
    "rtu / customer requests": {"equipmentType": ["RTU"]},
    "airside unit": {"equipmentType": ["RTU", "Split System", "AHU", "Heat Pump", "Water-Source Heat Pump", "Fan Coil", "VAV", "FPVAV"]},
    # equipment type groups
    "cooling equipment": {"equipmentType": COOLING},
    "condensing equipment": {"equipmentType": COOLING},
    "cooling test performed": {"equipmentType": COOLING},
    "refrigeration equipment": {"equipmentType": REFRIG},
    "compressor present": {"equipmentType": REFRIG},
    "cooling system can operate": {"equipmentType": REFRIG},
    "cooling system operating": {"equipmentType": REFRIG},
    "service valves/ports present": {"equipmentType": REFRIG},
    "service ports present": {"equipmentType": REFRIG},
    "valve packings present": {"equipmentType": REFRIG},
    "sight glass present": {"equipmentType": REFRIG},
    "sight glass / applicable compressor": {"equipmentType": REFRIG},
    "applicable metering device": {"equipmentType": REFRIG},
    "applicable metering device / system operating": {"equipmentType": REFRIG},
    "applicable system / system operating": {"equipmentType": REFRIG},
    "insulated lines present": {"equipmentType": REFRIG},
    "condenser fans present": {"equipmentType": CONDFAN},
    "condenser fans operating": {"equipmentType": CONDFAN},
    # heating types
    "gas heat": {"heatingType": ["Gas"]},
    "applicable gas heat": {"heatingType": ["Gas"]},
    "gas heat operating": {"heatingType": ["Gas"]},
    "fuel-fired heat": {"heatingType": FUEL_FIRED},
    "fuel-fired heat and instrument available": {"heatingType": FUEL_FIRED},
    "induced-draft heat": {"heatingType": FUEL_FIRED},
    "combustion analysis performed": {"heatingType": FUEL_FIRED},
    "electric heat": {"heatingType": ["Electric"]},
    "electric heat present and tested": {"heatingType": ["Electric"]},
    "heat pump": {"heatingType": ["Heat Pump"]},
    "heat pump with auxiliary heat": {"heatingType": ["Heat Pump"]},
    "hot-water heat": {"heatingType": ["Hot Water"]},
    "heating equipment": {"heatingType": HEATING_ANY},
    "heating mode can operate": {"heatingType": HEATING_ANY},
    "heating test performed": {"heatingType": HEATING_ANY},
    "applicable heating equipment": {"heatingType": HEATING_ANY},
    "hydronic/water-cooled equipment": HYDRONIC,
    "water-cooled equipment": HYDRONIC,
    "water-cooled condenser": HYDRONIC,
    "seasonal water coil": HYDRONIC,
    # mounting
    "rooftop or exterior units": {"mount": ROOF_MOUNTS},
    # prior answers
    "any access problem": {"whenCode": "T003", "whenIn": ["Issue", "Fail"]},                 # rule 3
    "full scope = no": {"whenCode": "F030", "whenIn": ["No"]},
    "call center notified": {"whenCode": "F032", "whenIn": ["Yes"]},
    "filters replaced": {"whenCode": "F076", "whenIn": ["Yes"]},
    "charge adjusted": {"whenCode": "F110", "whenIn": ["Low", "High"]},
    "refrigerant added": {"whenCode": "F111", "whenIn": ["Yes"]},
    "leak found": {"whenCode": "F113", "whenIn": ["Yes"]},
    "dirty grilles found": {"whenCode": "T074", "whenIn": ["Dirty", "Blocked"]},
    "static readings taken": {"whenCode": "F097", "whenIn": ["*"]},
    "pressure readings taken": {"whenCode": "F104", "whenIn": ["*"]},
    "compressor amp readings taken": {"whenCode": "F126", "whenIn": ["*"]},
    "motor amp readings taken": {"whenCode": "F129", "whenIn": ["*"]},
    "supply temperature recorded": {"whenCode": "F092", "whenIn": ["*"]},
    "performance readings taken": {"whenCode": "F092", "whenIn": ["*"]},
    "customer signoff required": {"whenCode": "T083", "whenIn": ["Completed"]},
}
TRIVIAL_TRIGGERS = {"every visit", "each unit", "every unit", "each unit / visit", "customer-specific", ""}

# Repeat keys (rule 9) and calculated fields
REPEAT_KEYS = {"F104": "Circuit", "F105": "Circuit", "F106": "Circuit", "F107": "Circuit", "F108": "Circuit", "F109": "Circuit",
               "F126": "Compressor", "F127": "Compressor", "F131": "HeatStage", "F168": "HeatStage", "F124": "Phase"}
CALC_FORMULAS = {"F095": "{F093}-{F092}", "F096": "{F092}-{F093}", "F099": "{F098}-{F097}", "F125": "100*MAXDEV({F124})/AVG({F124})"}
SEASON_BY_CODE = {"F095": SEASON_COOL, "F096": SEASON_HEAT, "F126": SEASON_COOL, "F127": SEASON_COOL, "F128": SEASON_COOL,
                  "F131": SEASON_HEAT, "T058": SEASON_HEAT, "T072": SEASON_HEAT}
HYDRONIC_CODES = {"F174", "F175", "F176", "F177", "F178", "F179", "F180", "T071", "T072"}   # in the heating section but not heating-season only
INACTIVE_CODES = {"F181": "Dictionary note: deficiency ID is just the quote WO#"}

# Named answer lists from the Dropdown Lists sheet, with overrides where the field's list is the superset
SHEET_LIST_OVERRIDES = {
    "Heating Type": HEATING_TYPES,
    "Equipment Type": EQUIPMENT_TYPES,
    "Severity": ["Observation", "Needs Service", "Urgent", "Safety/Critical", "Failed"],
    "PM Season": ["Spring Cooling", "Fall Heating", "Full PM", "Inspection Only", "Filter Change", "Inventory Only"],
}
FAIL_WORDS = ["fail", "needs", "replace", "leak", "issue", "damaged", "dirty", "worn", "cracked", "glazed", "frayed", "loose",
              "noisy", "out of range", "binding", "stuck", "alarm", "offline", "suspect", "wet", "discolored", "low", "high",
              "restricted", "corroded", "rubbing", "unsupported", "missing", "blocked", "disconnected", "hole", "rusted", "sticking",
              "weak", "uneven", "open element", "poor", "quote", "shut down", "intermittent", "severely", "fault", "calibration",
              "balance", "misaligned", "tight", "dry", "oil-contaminated", "partially", "declined", "unsafe", "hazard", "urgent", "safety",
              "communication", "imminent", "unit down", "reduced capacity", "comfort issue", "energy waste", "water leak", "not accessible"]
NOT_FAIL = {"no", "n/a", "not applicable", "none", "not needed", "not in scope", "not available", "no manager available",
            "unable", "unable to test", "unable to verify", "unable to fully inspect", "unknown", "monitor", "adjusted", "lubricated",
            "cleaned", "corrected", "tightened", "replaced", "already correct", "completed", "complete", "pass", "pass / completed",
            "good", "excellent", "fair", "proper", "normal", "clean", "aligned", "satisfied", "yes", "cooling", "heating", "fan only", "off",
            "lightly dirty", "occupied", "unoccupied", "custom", "24/7", "measured tension", "further diagnosis", "clean", "adjust", "repair"}

UNIT_PATTERNS = [(r"°F\s*/\s*%", "°F/%"), (r"°F", "°F"), (r"psig", "psig"), (r"in\. w\.c\.", "in. w.c."), (r"\bamps?\b", "A"),
                 (r"\bvolts?\b", "V"), (r"\bRPM\b", "RPM"), (r"\bppm\b", "ppm"), (r"\bpercent\b|%", "%"), (r"\bCFM\b", "CFM"),
                 (r"psi\s*/\s*gpm", "psi/gpm"), (r"[µμ]F", "µF"), (r"[µμ]A", "µA"), (r"\byears?\b", "yr"), (r"\btons?\b", "tons")]


# ---------------------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------------------
def s(v) -> str:
    if v is None:
        return ""
    return str(v).replace(" ", " ").strip()


def nstr(v) -> str:
    """T-SQL N'...' literal or NULL."""
    if v is None or v == "":
        return "NULL"
    return "N'" + str(v).replace("'", "''") + "'"


def num(v) -> str:
    return "NULL" if v is None else repr(round(float(v), 4))


def bit(v) -> str:
    return "NULL" if v is None else ("1" if v else "0")


def cond_json(c) -> str:
    return nstr(json.dumps(c, ensure_ascii=False, separators=(",", ":"))) if c else "NULL"


def split_values(raw: str) -> list[str]:
    raw = raw.replace("\n", " ")
    parts = raw.split(";") if ";" in raw else raw.split(",")
    out = []
    for p in parts:
        p = p.strip().rstrip(".")
        if not p:
            continue
        m = re.match(r"^(.*?),\s*(NA|N/A|Other.*)$", p)     # "Unable to inspect, NA", "York, Other (allow box for other)"
        if m and ";" in raw:
            out.append(m.group(1).strip())
            p = m.group(2)
        p = re.sub(r"\s*\(allow box for other\)", "", p).strip()
        if p and p not in out:
            out.append(p)
    return out


def fail_values(values: list[str]) -> list[str]:
    out = []
    for v in values:
        lv = v.lower()
        if lv in NOT_FAIL:
            continue
        if any(w in lv for w in FAIL_WORDS):
            out.append(v)
    return out


def parse_unit(allowed: str) -> str | None:
    for pat, unit in UNIT_PATTERNS:
        if re.search(pat, allowed, re.IGNORECASE):
            return unit
    return None


def is_whole_number(allowed: str) -> bool:
    return bool(re.search(r"whole number|years|quantity|sequential|^\d+(,\s*\d+)+", allowed, re.IGNORECASE))


# ---------------------------------------------------------------------------------------
# workbook readers
# ---------------------------------------------------------------------------------------
def read_sheet(wb, name, header_row):
    ws = wb[name]
    rows = list(ws.iter_rows(min_row=header_row, values_only=True))
    header = [s(c) for c in rows[0]]
    return header, [[s(c) for c in r] for r in rows[1:]]


def read_dictionary(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    _, fields = read_sheet(wb, "Master Data Fields", 3)
    _, tasks = read_sheet(wb, "PM Tasks", 3)
    _, photos = read_sheet(wb, "Photo Requirements", 3)
    _, lists = read_sheet(wb, "Dropdown Lists", 3)
    fields = [r for r in fields if r[1]]
    tasks = [r for r in tasks if r[0]]
    photos = [r for r in photos if r[0]]
    named = collections.OrderedDict()
    for r in lists:
        if r[0] and r[1]:
            named.setdefault(r[0], []).append(r[1])
    return fields, tasks, photos, named


def read_pricing(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    _, rows = read_sheet(wb, "Sheet1", 1)
    return [r for r in rows if r[0] or r[1]]


# ---------------------------------------------------------------------------------------
# question building
# ---------------------------------------------------------------------------------------
class Question:
    def __init__(self, **kw):
        self.code = kw["code"]
        self.section = kw["section"]
        self.clat = kw["clat"]
        self.question = kw["question"]
        self.requirement = kw["requirement"]
        self.condition = kw.get("condition")
        self.list_name = kw.get("list_name")
        self.values = kw.get("values")           # list[str] for inline lists
        self.datatype = kw.get("datatype")
        self.unit = kw.get("unit")
        self.repeatkey = kw.get("repeatkey")
        self.calc = kw.get("calc")
        self.writesto = kw.get("writesto", "Visit")
        self.photorequired = kw.get("photorequired", "No")
        self.phototiming = kw.get("phototiming")
        self.linkedcode = kw.get("linkedcode")
        self.estminutes = kw.get("estminutes")
        self.seasons = kw.get("seasons")
        self.triggernote = kw.get("triggernote")
        self.active = kw.get("active", True)
        self.source = kw.get("source", "")


class Report:
    def __init__(self):
        self.lines = []
        self.skipped = []
        self.unstructured = []
        self.notes = []

    def skip(self, code, reason):
        self.skipped.append((code, reason))

    def unstructured_trigger(self, code, trigger):
        self.unstructured.append((code, trigger))

    def note(self, text):
        self.notes.append(text)


def structure_trigger(code: str, trigger: str, report: Report):
    key = trigger.lower().strip()
    if key in TRIGGER_CONDITIONS:
        return json.loads(json.dumps(TRIGGER_CONDITIONS[key])), None
    if key in TRIVIAL_TRIGGERS:
        return None, None
    report.unstructured_trigger(code, trigger)
    return None, trigger


def build_field_questions(fields, report: Report, list_registry):
    out = []
    for r in fields:
        (_cl, fid, section, subsection, question, input_type, allowed, requirement, trigger, _src, photo, hours, notes) = (r + [""] * 13)[:13]
        secnum = int(section[0]) if section[:1].isdigit() else None
        if secnum in (None, 1, 2, 3):
            reason = {None: "Section ???? - office items (rule 1)", 1: "Section 1 - already built (rule 1)",
                      2: "Section 2 - becomes PMRule columns (rule 2)", 3: "Section 3 - ticket entry / location profile, not a tech question"}[secnum]
            report.skip(fid, reason)
            continue
        code = "F027" if fid == "F0TBD" else fid                                                    # rule 3
        section_name = FIELD_SECTION_MAP.get((secnum, subsection))
        if not section_name:
            raise SystemExit(f"{fid}: no section mapping for ({secnum}, {subsection})")
        phase = PHASE_BY_SECTION_NUMBER[secnum]

        it = input_type.lower()
        values = None
        list_name = None
        datatype = None
        clat = CLAT_TEXTBOX
        unit = None
        calc = None
        repeatkey = REPEAT_KEYS.get(code)
        if it in ("dropdown", "dropdown/text", "text/dropdown", "dropdown/number", "multi-select", "multi-select/text",
                  "yes/no", "yes/no/na", "yes/no/unknown", "repeating yes/no"):
            clat = CLAT_DROPDOWN
            values = split_values(allowed) if allowed else []
            if it == "repeating yes/no":
                values = ["Yes", "No"]
            datatype = "MultiSelect" if it.startswith("multi-select") else ("Bool" if it == "yes/no" else "Text")
        elif it in ("number", "repeating number"):
            datatype = "Number" if is_whole_number(allowed) else "Decimal"
            unit = "V" if code == "F124" else parse_unit(allowed)      # F124 "L1-L2, L2-L3, L1-L3 or L-N" names no unit
        elif it == "formula/number":
            clat = CLAT_READONLY
            datatype = "Decimal"
            unit = parse_unit(allowed)
            calc = CALC_FORMULAS.get(code)
            if not calc:
                report.note(f"{code}: Formula/Number without a known formula")
        elif it in ("number/text", "text/number"):
            datatype = "Text"
            unit = parse_unit(allowed)
        elif it == "date":
            datatype = "Date"
        elif it in ("date-time", "date-time/text", "time"):
            datatype = "DateTime"
        elif it == "signature":
            clat = CLAT_SIGNATURE
        elif it == "image/signature":
            clat = CLAT_PHOTO
        elif it == "auto-number":
            clat = CLAT_READONLY
            datatype = "Number"
        else:  # text, long text, phone, email, dimensions, date/number
            datatype = "Text"

        if values is not None:
            list_name = list_registry.register(code, values)
            if list_name:
                values = None

        if code == "F049":
            list_name, values = "Equipment Type", None                                             # rule 13 vocabulary
        if code == "F059":
            list_name, values = "Heating Type", None
        if code == "F184":
            list_name, values = "Severity", None
        if code == "F067":
            list_name, values = "Condition", None
        if code == "F068":
            list_name, values = "Overall Recommendation", None

        req = requirement
        if req not in ("Always", "Recommended", "Optional", "Conditional", "Configured"):
            raise SystemExit(f"{fid}: unknown requirement {requirement}")
        condition, triggernote = structure_trigger(code, trigger, report)
        if code == "F051-1":
            condition, triggernote = {"equipmentType": ["VRF"]}, None
        elif code == "F051-2":
            condition, triggernote = {"equipmentType": ["Chiller"]}, None
        elif code == "F051-3":
            condition, triggernote = {"equipmentType": ["Other"]}, None
        elif code == "F051-4":
            condition, triggernote = None, "PTAC units (no PTAC equipment type in the list)"
        if code in ("F051-1", "F051-2", "F051-3", "F051-4"):
            req = "Conditional"
        if code in FORCE_CONFIGURED:
            req = "Configured"                                                                       # rules 4 and 7
        if condition and req not in ("Conditional", "Configured"):
            req = "Conditional"

        pl = photo.lower()
        photorequired = "Recommended" if pl in ("y", "maybe?") else "No"
        if pl == "maybe?":
            triggernote = (triggernote + " | " if triggernote else "") + "photo: Maybe?"
        est = round(float(hours) * 60, 2) if hours else None
        seasons = SEASON_BY_CODE.get(code)
        if section_name == "Cooling & Refrigeration Readings":
            seasons = SEASON_COOL
        elif section_name == "Heating & Hydronic Readings" and code not in HYDRONIC_CODES:
            seasons = SEASON_HEAT
        out.append(Question(code=code, section=section_name, clat=clat, question=question, requirement=req, condition=condition,
                            list_name=list_name, values=values, datatype=datatype, unit=unit, repeatkey=repeatkey, calc=calc,
                            writesto="Asset" if phase == "Asset" else "Visit", photorequired=photorequired, estminutes=est,
                            seasons=seasons, triggernote=triggernote, active=code not in INACTIVE_CODES, source="field"))
    return out


TASK_PHOTO_MAP = {
    "required": ("Always", None), "if issue": ("IfIssue", None), "issue only": ("IfIssue", None), "issue photos": ("IfIssue", None),
    "optional": ("Recommended", None), "instrument display optional": ("Recommended", "During"),
    "instrument displays recommended": ("Recommended", "During"), "instrument display": ("Recommended", "During"),
    "interior photo": ("Recommended", "During"), "before/after": ("Always", "BeforeAfter"), "before and after": ("Always", "BeforeAfter"),
    "before photo": ("Always", "Before"), "after photo": ("Always", "After"), "before/after if dirty": ("IfIssue", "BeforeAfter"),
    "required if dirty": ("IfIssue", None), "required if yes": ("IfIssue", None),
    "after photo when required": ("CustomerConfigured", "After"), "required when customer requests": ("CustomerConfigured", None),
    "no": ("No", None),
}


def build_task_questions(tasks, report: Report, list_registry):
    out = []
    for r in tasks:
        (tid, area, task, response, reading, applies, requirement, photo, _src, _notes) = (r + [""] * 10)[:10]
        section_name = TASK_AREA_MAP.get(area)
        if not section_name:
            raise SystemExit(f"{tid}: no section mapping for area {area}")
        if response.lower() == "condition status":
            list_name, values = "Condition", None
        else:
            values = [v.strip() for v in response.split(" / ") if v.strip()]
            list_name = list_registry.register(tid, values)
            if list_name:
                values = None
        req = {"Always": "Always", "Recommended": "Recommended", "Conditional": "Conditional", "Customer-configured": "Configured"}[requirement]
        condition, triggernote = structure_trigger(tid, applies, report)
        if tid in FORCE_CONFIGURED:
            req = "Configured"
        if condition and req not in ("Conditional", "Configured"):
            req = "Conditional"
        if tid == "T014":
            condition, triggernote = {"whenCode": "T013", "whenIn": ["Completed"]}, None
        photorequired, timing = TASK_PHOTO_MAP[photo.lower()]
        question = task if not reading else f"{task} (record: {reading})"
        seasons = SEASON_BY_CODE.get(tid)
        if section_name == "Cooling & Refrigeration Readings":
            seasons = SEASON_COOL
        elif section_name == "Heating & Hydronic Readings" and tid not in HYDRONIC_CODES:
            seasons = SEASON_HEAT
        out.append(Question(code=tid, section=section_name, clat=CLAT_DROPDOWN, question=question, requirement=req, condition=condition,
                            list_name=list_name, values=values, datatype="Text", photorequired=photorequired, phototiming=timing,
                            linkedcode=tid if photorequired == "IfIssue" else None, seasons=seasons, triggernote=triggernote, source="task"))
    return out


def build_photo_questions(photos, report: Report):
    out = []
    for r in photos:
        (pid, evidence, frequency, when, trigger, _stamp, _src, purpose) = (r + [""] * 8)[:8]
        section_name, repeatkey = PHOTO_SECTION_MAP[pid]
        photorequired, req, condition, linked = PHOTO_TRIGGER_MAP[pid]
        timing = PHOTO_TIMING_MAP.get(when.lower())
        if timing is None:
            report.note(f"{pid}: unknown 'When Taken' value {when!r}")
        seasons = SEASON_COOL if section_name == "Cooling & Refrigeration Readings" else (SEASON_HEAT if section_name == "Heating & Hydronic Readings" else None)
        out.append(Question(code=pid, section=section_name, clat=CLAT_PHOTO, question=f"Photo: {evidence}", requirement=req, condition=condition,
                            repeatkey=repeatkey, photorequired=photorequired, phototiming=timing, linkedcode=linked, seasons=seasons,
                            triggernote=f"{trigger} ({frequency})", source="photo"))
    return out


class ListRegistry:
    """Rule 10: value sets used by two or more questions become named FormAnswerList rows."""

    def __init__(self, named):
        self.sheet_lists = collections.OrderedDict()
        for name, values in named.items():
            self.sheet_lists[name] = SHEET_LIST_OVERRIDES.get(name, values)
        self.sheet_lists.setdefault("PM Season", SHEET_LIST_OVERRIDES["PM Season"])
        self.usage = collections.OrderedDict()   # key -> [codes]

    @staticmethod
    def key(values):
        return ";".join(values)

    def register(self, code, values):
        self.usage.setdefault(self.key(values), []).append(code)
        return None   # resolved in finalize()

    def finalize(self, questions):
        names = {}
        for key, codes in self.usage.items():
            if len(codes) < 2:
                continue
            values = key.split(";")
            lk = [v.lower() for v in values]
            if lk == ["yes", "no", "n/a"]:
                name = "Yes/No/NA"
            elif lk == ["yes", "no"]:
                name = "Yes/No"
            elif lk == ["yes", "no", "unknown"]:
                name = "Yes/No/Unknown"
            elif lk == ["pass", "fail", "n/a"]:
                name = "Pass/Fail/NA"
            else:
                name = "Task: " + " / ".join(values)
                if len(name) > 60:
                    name = name[:57] + "..."
            names[key] = name
        generated = collections.OrderedDict()
        for q in questions:
            if q.values is not None:
                name = names.get(self.key(q.values))
                if name:
                    generated[name] = q.values
                    q.list_name, q.values = name, None
        lists = collections.OrderedDict(self.sheet_lists)
        lists.update(generated)
        return lists


# ---------------------------------------------------------------------------------------
# SQL writers
# ---------------------------------------------------------------------------------------
def header(title, source_files):
    return (f"-- {title}\n-- Generated by sql/forms/generate_form_seed.py on {TODAY} from {', '.join(source_files)}.\n"
            "-- Do not hand-edit: change the workbook or the generator and regenerate.\n"
            "-- Requires create_form_tables.sql. Safe to re-run (see the notes in each block).\n\n")


def write_template_sql(path: Path, lists, questions, report: Report):
    L = [header("Form template seed: HVAC PM (answer lists, sections, questions)", ["Universal_HVAC_PM_Master_Checklist_Data_Dictionary.xlsx"])]
    L.append("SET NOCOUNT ON;\nSET XACT_ABORT ON;\n")
    L.append(f"DECLARE @t_id INT = (SELECT TOP 1 t_id FROM dbo.Trade WHERE t_trade = N'{PARENT_TRADE}' AND t_id_parent IS NULL);\n")
    L.append("IF @t_id IS NULL BEGIN RAISERROR('Parent trade HVAC not found', 16, 1); RETURN; END;\n\n")

    L.append("-- Answer lists: inserted when missing, existing rows are left untouched (admin edits survive a re-run).\n")
    for name, values in lists.items():
        fails = fail_values(values)
        L.append(f"IF NOT EXISTS (SELECT 1 FROM dbo.FormAnswerList WHERE fal_name = {nstr(name)})\n"
                 f"    INSERT INTO dbo.FormAnswerList (fal_name, fal_values, fal_failvalues) VALUES ({nstr(name)}, {nstr(';'.join(values))}, {nstr(';'.join(fails)) if fails else 'NULL'});\n")

    L.append("\n-- Template: version 1 is replaced in place while nothing references it; once form rules point at it,\n"
             "-- a re-run seeds the next version and deactivates the previous one (rules keep their pinned version).\n")
    L.append("DECLARE @version INT = 1;\nDECLARE @ft_id INT;\n")
    L.append("DECLARE @existing INT = (SELECT ft_id FROM dbo.FormTemplate WHERE t_id = @t_id AND ft_version = @version);\n")
    L.append("IF @existing IS NOT NULL\nBEGIN\n"
             "    IF EXISTS (SELECT 1 FROM dbo.FormRule WHERE ft_id = @existing)\n"
             "       OR EXISTS (SELECT 1 FROM dbo.xrefFormRuleQuestion x JOIN dbo.FormQuestion q ON q.fq_id = x.fq_id\n"
             "                  JOIN dbo.FormSection s ON s.fs_id = q.fs_id WHERE s.ft_id = @existing)\n"
             "    BEGIN\n"
             "        SET @version = (SELECT MAX(ft_version) + 1 FROM dbo.FormTemplate WHERE t_id = @t_id);\n"
             "        PRINT 'HVAC PM v1 is referenced by form rules; seeding as version ' + CAST(@version AS VARCHAR(10));\n"
             "        UPDATE dbo.FormTemplate SET ft_active = 0, ft_modifieddatetime = GETDATE() WHERE t_id = @t_id AND ft_active = 1;\n"
             "    END\n"
             "    ELSE\n"
             "    BEGIN\n"
             "        PRINT 'Replacing unreferenced HVAC PM v1';\n"
             "        DELETE q FROM dbo.FormQuestion q JOIN dbo.FormSection s ON s.fs_id = q.fs_id WHERE s.ft_id = @existing;\n"
             "        DELETE FROM dbo.FormSection WHERE ft_id = @existing;\n"
             "        DELETE FROM dbo.FormTemplate WHERE ft_id = @existing;\n"
             "    END\n"
             "END;\n")
    L.append("BEGIN TRANSACTION;\n")
    note = f"Seeded from Universal_HVAC_PM_Master_Checklist_Data_Dictionary.xlsx by generate_form_seed.py on {TODAY}"
    L.append(f"INSERT INTO dbo.FormTemplate (t_id, ft_name, ft_version, ft_active, ft_note) VALUES (@t_id, {nstr(TEMPLATE_NAME)}, @version, 1, {nstr(note)});\n"
             "SET @ft_id = SCOPE_IDENTITY();\n\n")

    L.append("-- Sections\n")
    for i, (name, phase, repeat, cond) in enumerate(SECTIONS, start=1):
        L.append(f"INSERT INTO dbo.FormSection (ft_id, fs_name, fs_phase, fs_order, fs_repeatperunit, fs_condition)\n"
                 f"    VALUES (@ft_id, {nstr(name)}, {nstr(phase)}, {i}, {repeat}, {cond_json(cond)});\n"
                 f"DECLARE @s{i} INT = SCOPE_IDENTITY();\n")

    L.append("\n-- Questions (tasks, then fields, then photos within each section; dictionary order inside each group)\n")
    order_in_section = collections.Counter()
    src_rank = {"task": 0, "field": 1, "photo": 2}
    questions_sorted = sorted(questions, key=lambda q: (SECTION_INDEX[q.section], src_rank[q.source]))
    for q in questions_sorted:
        order_in_section[q.section] += 1
        sec_var = f"@s{SECTION_INDEX[q.section] + 1}"
        fal = f"(SELECT fal_id FROM dbo.FormAnswerList WHERE fal_name = {nstr(q.list_name)})" if q.list_name else "NULL"
        L.append("INSERT INTO dbo.FormQuestion (fs_id, fq_code, clat_id, fq_question, fq_order, fq_requirement, fq_condition, fal_id, fq_answervalues, "
                 "fq_datatype, fq_unit, fq_repeatkey, fq_calcformula, fq_writesto, fq_photorequired, fq_phototiming, fq_linkedcode, fq_estminutes, "
                 "fq_seasons, fq_triggernote, fq_active)\n"
                 f"    SELECT {sec_var}, {nstr(q.code)}, {q.clat}, {nstr(q.question)}, {order_in_section[q.section]}, {nstr(q.requirement)}, {cond_json(q.condition)}, {fal}, "
                 f"{nstr(';'.join(q.values)) if q.values else 'NULL'}, {nstr(q.datatype)}, {nstr(q.unit)}, {nstr(q.repeatkey)}, {nstr(q.calc)}, {nstr(q.writesto)}, "
                 f"{nstr(q.photorequired)}, {nstr(q.phototiming)}, {nstr(q.linkedcode)}, {num(q.estminutes)}, {nstr(q.seasons)}, {nstr(q.triggernote)}, {1 if q.active else 0};\n")
    L.append("\nCOMMIT TRANSACTION;\n")
    L.append("DECLARE @sections INT = (SELECT COUNT(*) FROM dbo.FormSection WHERE ft_id = @ft_id);\n"
             "DECLARE @questions INT = (SELECT COUNT(*) FROM dbo.FormQuestion q JOIN dbo.FormSection s ON s.fs_id = q.fs_id WHERE s.ft_id = @ft_id);\n"
             "PRINT 'HVAC PM template seeded: version ' + CAST(@version AS VARCHAR(10)) + ', ' + CAST(@sections AS VARCHAR(10)) + ' sections, ' + CAST(@questions AS VARCHAR(10)) + ' questions';\n")
    path.write_text("".join(L), encoding="utf-8")


def write_asset_categories_sql(path: Path):
    L = [header("Seed: the dictionary's 12 HVAC equipment types as AssetCategory rows (rule 13)", ["Universal_HVAC_PM_Master_Checklist_Data_Dictionary.xlsx (Dropdown Lists)"])]
    L.append("-- Existing HVAC categories (Chiller, Energy Management System, Exhaust Fan, Hood Vent, HVAC, Thermostat) are left untouched.\n"
             "-- Required-field flags follow the existing HVAC rows (manufacturer, model, serial, description all required).\n")
    L.append("SET NOCOUNT ON;\n")
    L.append(f"DECLARE @t_id INT = (SELECT TOP 1 t_id FROM dbo.Trade WHERE t_trade = N'{PARENT_TRADE}' AND t_id_parent IS NULL);\n")
    L.append("IF @t_id IS NULL BEGIN RAISERROR('Parent trade HVAC not found', 16, 1); RETURN; END;\n")
    L.append("DECLARE @inserted INT = 0;\n")
    for name in EQUIPMENT_TYPES:
        L.append(f"IF NOT EXISTS (SELECT 1 FROM dbo.AssetCategory WHERE t_id = @t_id AND LOWER(LTRIM(RTRIM(asc_category))) = LOWER({nstr(name)}))\n"
                 f"BEGIN\n    INSERT INTO dbo.AssetCategory (t_id, asc_category, asc_manufacturer, asc_modelnumber, asc_serialnumber, asc_description)\n"
                 f"        VALUES (@t_id, {nstr(name)}, 1, 1, 1, 1);\n    SET @inserted += 1;\nEND;\n")
    L.append("PRINT 'HVAC asset categories inserted: ' + CAST(@inserted AS VARCHAR(10));\n")
    path.write_text("".join(L), encoding="utf-8")


# Pricing sheet mapping (rule 14 + SEED_DECISIONS.md)
CALLCENTER_ALIASES = {
    "Boss": "Boss Facility Services Inc",
    "Brinco": "Brinco - National HVAC Management Services",
    "Broadway": "Broadway National/Servco",
    "CLS": "CLS Facility Services",
    "Direct Commercial": "Direct Commercial",
    "Legacy": "Legacy Group Enterprises Inc",
    "Powerhouse Retail": "Powerhouse Retail Services",
    "Retail Mechanical": "Retail Mechanical Services",
    "SMS": "SMS Assist",
    "Total Comfort Group": "Total Comfort Group",
    "Vixxo": "First Service Networks/Vixxo",
}
COMPANY_ALIASES = {
    ("Boss", "Drivetime"): "Drive Time",
    ("Brinco", "Burlington"): "Burlington Coat Factory",
    ("CLS", "Jockey"): "Jockey International",
    ("Direct Commercial", "First Cash"): "FirstCash, Inc.",
    ("Direct Commercial", "Jeni's Splendid Ice Cream"): "Jeni's Splendid Ice Creams",
    ("SMS", "Signet Jewlers"): "Signet Jewelers",
    ("Legacy", "XPO"): "XPO LOGISTICS",
}
NTE_MAP = {
    'your "final" total due should match the set nte': "MatchSetNte",
    'nte in portal should match your "final" total due, if not reach out for increase': "MatchPortalNte",
    "do not follow nte, no portal to match": "IgnoreNte",
}
FIRST_TIME_MAP = {
    "can bill a trip charge in addition to contract pricing": "TripChargeExtra",
    "can bill a trip charge in addition to set nte cost": "TripChargeExtra",
    "bill trip, labor, all materials": "BillTripLaborMaterials",
    "no": "None",
}


def map_customer_form(text: str):
    t = text.lower().replace("\n", " ").strip()
    if not t:
        return None, None, None
    if t.startswith("no, use evolution"):
        return "None", None, None
    if "airtable" in t:
        return "ExternalLink", "Airtable link (URL to be added on the rules tab)", "Portal"
    if t.startswith("go canvas"):
        return "ExternalLink", "Go Canvas (URL to be added on the rules tab)", "Portal"
    if "sms portal" in t:
        return "Portal", "SMS portal under Serviced Assets; fill out the data sheet only if the portal is not working", "Portal"
    if "2 chklist required" in t:
        return "CustomerPdfForm", "Provided by CLS; 2 checklists required", None
    m = re.match(r"yes,?\s*(.*?provided by\s+.+)$", t)
    if m:
        return "CustomerPdfForm", m.group(1).strip().capitalize(), None
    return "CustomerPdfForm", text.strip(), None


def parse_inclusions(comment: str):
    c = comment.lower()
    inc = {"filtersincluded": None, "nofilterchange": None, "freefilters": None, "freebelts": None, "beltschargeable": None, "coilcleaner": None}
    if "filters included in pm cost" in c or "filters are included in pm cost" in c:
        inc["filtersincluded"] = True
    m = re.search(r"(\d+)\s+filters?\s+free", c)
    if m:
        inc["freefilters"] = int(m.group(1))
        inc["filtersincluded"] = True
    m = re.search(r"(\d+)\s+belt change", c)
    if m:
        inc["freebelts"] = int(m.group(1))
    if "no filters should be changed or charged" in c:
        inc["nofilterchange"] = True
        inc["filtersincluded"] = False
    if "coil cleaner can be charged" in c or "belt cost" in c:
        inc["beltschargeable"] = True
    m = re.search(r"coil cleaner \$([\d.]+)", c)
    if m:
        inc["coilcleaner"] = float(m.group(1))
    return inc


def write_rules_sql(path: Path, pricing_rows, report: Report):
    L = [header("Seed: form rules + PM terms from the PM Pricing Info sheet (rule 14)", ["Copy of PM Pricing Info.xlsx"])]
    L.append("-- One FormRule + PMRule per sheet row, matched case-insensitively on call center + company name (aliases in the generator).\n"
             "-- Rows whose pairing does not exist are PRINTed and skipped; rows that already have a rule are left untouched.\n"
             "-- Blank pricing rows still get a rule (billing mode NULL) so the tab has somewhere to put the team's answers.\n")
    L.append("SET NOCOUNT ON;\n")
    L.append(f"DECLARE @t_id INT = (SELECT TOP 1 t_id FROM dbo.Trade WHERE t_trade = N'{PARENT_TRADE}' AND t_id_parent IS NULL);\n")
    L.append("DECLARE @ft_id INT = (SELECT TOP 1 ft_id FROM dbo.FormTemplate WHERE t_id = @t_id AND ft_active = 1 ORDER BY ft_version DESC);\n")
    L.append("IF @t_id IS NULL OR @ft_id IS NULL BEGIN RAISERROR('HVAC trade or active HVAC PM template not found - run seed_form_template_hvac.sql first', 16, 1); RETURN; END;\n")
    L.append("DECLARE @pmbm_setnte INT = (SELECT pmbm_id FROM dbo.PMBillingMode WHERE pmbm_mode = 'SetNte');\n"
             "DECLARE @pmbm_tiered INT = (SELECT pmbm_id FROM dbo.PMBillingMode WHERE pmbm_mode = 'ContractTiered');\n"
             "DECLARE @xccc_id INT, @fr_id INT, @pmr_id INT, @inserted INT = 0, @skipped INT = 0, @existing INT = 0;\n\n")
    for r in pricing_rows:
        (cc, company, contract, attachment, nte, first_time, checklist, comments, contact, per_year) = (r + [""] * 10)[:10]
        cc_name = CALLCENTER_ALIASES.get(cc)
        label = f"{cc} / {company}".replace("'", "''")
        if not cc_name:
            report.note(f"pricing: call center '{cc}' has no CallCenter row - skipped {company}")
            L.append(f"PRINT 'SKIP (no call center): {label}'; SET @skipped += 1;\n")
            continue
        company_name = COMPANY_ALIASES.get((cc, company), company)
        billing = None
        if contract.strip().upper() == "YES":
            billing = "@pmbm_tiered"
        elif contract.strip().upper() == "NO":
            billing = "@pmbm_setnte"
        nte_code = NTE_MAP.get(nte.lower().strip()) if nte else None
        if nte and not nte_code:
            report.note(f"pricing: unknown NTE guideline text for {cc}/{company}: {nte}")
        ft_code = FIRST_TIME_MAP.get(first_time.lower().strip()) if first_time else None
        if first_time and not ft_code:
            report.note(f"pricing: unknown first-time rule text for {cc}/{company}: {first_time}")
        form, form_note, destination = map_customer_form(checklist)
        inc = parse_inclusions(comments)
        contact_name = contact_email = None
        if contact:
            parts = contact.split(":", 1)
            contact_name = parts[0].strip()
            contact_email = parts[1].strip() if len(parts) > 1 else None
        notes = [f"Seeded from PM Pricing Info sheet ({TODAY})"]
        if attachment:
            notes.append(f"Pricing contract attachment in Evo: {attachment.strip()}")
        if "same as hibbett" in comments.lower():
            notes.append("Pricing: same as Hibbett")
        enable_codes = []
        if inc["filtersincluded"]:
            enable_codes += CONFIGURED_BY_RULE["filtersincluded"]

        L.append(f"-- {cc} / {company}\n")
        L.append("SET @xccc_id = (SELECT TOP 1 x.xccc_id FROM dbo.xrefCompanyCallCenter x JOIN dbo.CallCenter cc ON cc.cc_id = x.cc_id JOIN dbo.Company c ON c.c_id = x.c_id\n"
                 f"                WHERE LOWER(LTRIM(RTRIM(cc.cc_name))) = LOWER({nstr(cc_name)}) AND LOWER(LTRIM(RTRIM(c.c_name))) = LOWER({nstr(company_name)}) ORDER BY x.xccc_active DESC, x.xccc_id);\n")
        L.append(f"IF @xccc_id IS NULL BEGIN PRINT 'SKIP (no company pairing): {label}'; SET @skipped += 1; END\n"
                 f"ELSE IF EXISTS (SELECT 1 FROM dbo.FormRule WHERE xccc_id = @xccc_id AND t_id = @t_id) BEGIN SET @existing += 1; END\n"
                 "ELSE BEGIN\n")
        L.append(f"    INSERT INTO dbo.FormRule (xccc_id, t_id, ft_id, fr_note) VALUES (@xccc_id, @t_id, @ft_id, {nstr(' | '.join(notes))});\n"
                 "    SET @fr_id = SCOPE_IDENTITY();\n")
        L.append("    INSERT INTO dbo.PMRule (fr_id, pmbm_id, pmr_nteguideline, pmr_firsttimerule, pmr_pricingnote, pmr_filtersincluded, pmr_nofilterchange, "
                 "pmr_freefilters, pmr_freebeltchangesperyear, pmr_beltschargeable, pmr_customerform, pmr_customerformnote, pmr_submissiondestination, "
                 "pmr_contactname, pmr_contactemail)\n"
                 f"        VALUES (@fr_id, {billing or 'NULL'}, {nstr(nte_code)}, {nstr(ft_code)}, {nstr(comments.strip() or None)}, {bit(inc['filtersincluded'])}, {bit(inc['nofilterchange'])}, "
                 f"{inc['freefilters'] if inc['freefilters'] is not None else 'NULL'}, {inc['freebelts'] if inc['freebelts'] is not None else 'NULL'}, {bit(inc['beltschargeable'])}, "
                 f"{nstr(form)}, {nstr(form_note)}, {nstr(destination)}, {nstr(contact_name)}, {nstr(contact_email)});\n"
                 "    SET @pmr_id = SCOPE_IDENTITY();\n")
        if inc["coilcleaner"] is not None:
            L.append(f"    INSERT INTO dbo.PMRateTier (pmr_id, pmrt_label, pmrt_firstunitprice, pmrt_addon, pmrt_addonlabel, pmrt_order) VALUES (@pmr_id, N'Coil cleaner', {inc['coilcleaner']:.2f}, 1, N'Coil cleaner', 1);\n")
        if enable_codes:
            codes = ", ".join(nstr(c) for c in enable_codes)
            L.append("    INSERT INTO dbo.xrefFormRuleQuestion (fr_id, fq_id, xfrq_enabled)\n"
                     f"        SELECT @fr_id, q.fq_id, 1 FROM dbo.FormQuestion q JOIN dbo.FormSection s ON s.fs_id = q.fs_id WHERE s.ft_id = @ft_id AND q.fq_code IN ({codes});\n")
        L.append("    SET @inserted += 1;\nEND;\n\n")
    L.append("PRINT 'Form rules (PM): inserted ' + CAST(@inserted AS VARCHAR(10)) + ', already present ' + CAST(@existing AS VARCHAR(10)) + ', skipped ' + CAST(@skipped AS VARCHAR(10));\n")
    path.write_text("".join(L), encoding="utf-8")


# ---------------------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dict", type=Path, default=DEFAULT_DICT)
    ap.add_argument("--pricing", type=Path, default=DEFAULT_PRICING)
    ap.add_argument("--out", type=Path, default=HERE)
    args = ap.parse_args()

    report = Report()
    fields, tasks, photos, named = read_dictionary(args.dict)
    registry = ListRegistry(named)
    questions = build_field_questions(fields, report, registry) + build_task_questions(tasks, report, registry) + build_photo_questions(photos, report)
    lists = registry.finalize(questions)

    codes = collections.Counter(q.code for q in questions)
    dupes = [c for c, n in codes.items() if n > 1]
    if dupes:
        raise SystemExit(f"duplicate question codes: {dupes}")
    for q in questions:
        if q.linkedcode and q.linkedcode not in codes:
            raise SystemExit(f"{q.code}: linked code {q.linkedcode} is not seeded")
        if q.condition and "whenCode" in q.condition and q.condition["whenCode"] not in codes:
            raise SystemExit(f"{q.code}: condition references {q.condition['whenCode']} which is not seeded")

    args.out.mkdir(parents=True, exist_ok=True)
    write_template_sql(args.out / "seed_form_template_hvac.sql", lists, questions, report)
    write_asset_categories_sql(args.out / "seed_asset_categories_hvac.sql")
    pricing_rows = read_pricing(args.pricing)
    write_rules_sql(args.out / "seed_form_rules_pm_hvac.sql", pricing_rows, report)

    # ---- report
    R = [f"Form seed generator report - {TODAY}", "=" * 60, ""]
    R.append(f"Dictionary rows: {len(fields)} fields, {len(tasks)} tasks, {len(photos)} photos")
    by_src = collections.Counter(q.source for q in questions)
    R.append(f"Questions seeded: {len(questions)} ({by_src['field']} fields, {by_src['task']} tasks, {by_src['photo']} photos); inactive: {[q.code for q in questions if not q.active]}")
    R.append(f"Answer lists: {len(lists)} ({len(registry.sheet_lists)} from the Dropdown Lists sheet, {len(lists) - len(registry.sheet_lists)} generated from shared value sets)")
    R.append("")
    R.append("Questions per section (phase):")
    per_sec = collections.Counter(q.section for q in questions)
    for name, phase, *_ in SECTIONS:
        R.append(f"  {per_sec[name]:4d}  {name} ({phase})")
    R.append("")
    R.append("Requirement mix: " + ", ".join(f"{k} {v}" for k, v in sorted(collections.Counter(q.requirement for q in questions).items())))
    R.append("Photo rule mix:  " + ", ".join(f"{k} {v}" for k, v in sorted(collections.Counter(q.photorequired for q in questions).items())))
    R.append(f"Structured conditions: {sum(1 for q in questions if q.condition)}; trigger kept as note only: {len(report.unstructured)}")
    R.append("")
    R.append(f"Skipped dictionary fields ({len(report.skipped)}):")
    for code, reason in report.skipped:
        R.append(f"  {code:8} {reason}")
    R.append("")
    R.append("Triggers not turned into a structured condition (kept in fq_triggernote):")
    for code, trig in report.unstructured:
        R.append(f"  {code:8} {trig}")
    R.append("")
    R.append("Configured questions enabled by PM rule columns:")
    for k, v in CONFIGURED_BY_RULE.items():
        R.append(f"  {k:24} {', '.join(v)}")
    R.append("")
    R.append("Answer lists:")
    for name, values in lists.items():
        R.append(f"  {name}: {'; '.join(values)}  [fail: {'; '.join(fail_values(values)) or '-'}]")
    R.append("")
    R.append(f"Pricing sheet rows: {len(pricing_rows)}; call-center aliases: {CALLCENTER_ALIASES}; company aliases: {COMPANY_ALIASES}")
    R.append("")
    R.append("Generator notes:")
    for n in report.notes:
        R.append(f"  {n}")
    (args.out / "seed_form_report.txt").write_text("\n".join(R) + "\n", encoding="utf-8")
    print("\n".join(R))


if __name__ == "__main__":
    main()
